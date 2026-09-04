"""레지스터 덤프 저장 / 로드 / 비교 CLI.

Sivers 가 준 레퍼런스 덤프(reference/Data_260729_DoosanSTMP_RegDump.xlsx)와
우리 bring-up 결과를 필드 단위로 대조한다. 판정 기준은 설계 스펙 8.5 참고:
잔차가 다이별 eFuse bias 에만 남으면 SW 포팅은 맞다.

실행:
  python -m cloudchaser.regdump --save out\\regs.csv
  python -m cloudchaser.regdump --diff reference\\Data_260729_DoosanSTMP_RegDump.xlsx
  python -m cloudchaser.regdump --diff out\\regs_efuse.csv --ours out\\regs_casper.csv
  python -m cloudchaser.regdump --fake --diff tests\\data\\golden_regs_tx_v1_split.csv

주의: 사용자에게 보이는 출력/예외는 영어로 작성한다(콘솔 한글 깨짐 방지).

xlsx 로더 참고 (Task 11 조사, reference/*.xlsx 실물을 직접 열어 확인):
실제 Sivers 파일은 시트 1개, 데이터가 열 B(주소)/C(값)에 있고 열 A는 항상 빈다.
헤더는 'Register Adress'(오타)/'Value'. 주소 열은 세 형식이 섞여 있다 -- 헥사
자릿수를 그대로 십진수로 오인식한 int(예: 1004 -> 실주소 0x1004), 'XXXX' 헥사
문자열(0x 접두 없음), '0xXXXX' 문자열. 게다가 Excel 이 '10E0'~'10E9' 10개 셀을
지수표기 숫자(10, 100, ... 1e10)로 자동변환해버려 원본 주소를 복구할 수 없는
행이 섞여 있다 -- 그 값을 그대로 흡수하면 이미 읽은 다른 정상 주소(예: 0x1000)
와 충돌해 값을 덮어쓰는 사고가 난다. 그래서 주소가 직전 주소보다 작거나
너무 멀리 뛰면(_XLSX_MAX_ADDR_GAP) 복구 불가능한 행으로 보고 건너뛴다.

Ruling 34: 건너뛴 행을 조용히 버리면 안 된다. "N differing / M compared"
한 줄만 보면 마치 전 구간을 다 비교한 것처럼 읽히는데, 실제로는 그 파싱
불가 행들이 애초에 비교 대상에 들어가지도 않았기 때문이다(이 마이그레이션
전체가 없애려는 것과 같은 모양의 '조용한 구멍'). 그래서 load_dump/​
_load_xlsx 는 dict 대신 RegDump(dict 의 부가정보 달린 서브클래스)를 반환하고,
건너뛴 행이 하나라도 있으면 로드 시점에 note 를 찍는다 -- CLI 로 실행하면
theirs/ours 로드 직후, 즉 diff 결과("N differing / M compared") 바로 위에
뜨므로 둘을 같이 읽게 된다. 아무것도 안 건너뛰었으면 조용하다.

Ruling 35 -- 주의: _addr_from_xlsx_cell 의 "int 셀 = 헥사 자릿수를 그대로
십진수로 적은 것" 규칙은 *이* Sivers 파일에서 실측으로 확인한 규칙이지,
xlsx 포맷 자체의 보편 규칙이 아니다. 다른 파일이 "평범한" 관례로 만들어졌다면
(즉 셀 4096 이 정말 정수 4096 = 0x1000 을 뜻한다면) 이 코드는 그걸 조용히
0x4096 으로 잘못 읽는다 -- 주소가 전부 틀렸는데 에러도 없이 "N differing / M
compared" 한 줄이 그럴듯하게 찍힌다. 이게 바로 이 마이그레이션 전체가
없애려는 실패 형태라서, 로더가 자기 자신의 가정을 못 믿는다는 걸 알려야
한다. 그래서 주소를 다 읽은 뒤 알려진 레지스터 범위(_KNOWN_ADDR_RANGES)
밖으로 떨어진 비율을 확인해서, 크면 경고만 찍는다 -- 자동으로 다른 해석으로
바꿔치기하지 않는다(잘못된 추측을 조용히 적용하는 것도 똑같이 위험하므로).

이 plausibility 체크는 magnitude(크기)만 본다는 한계가 있다 -- 주소가 전부
작은 값(예: 0..255)인 plain-decimal 파일은, 헥사 자릿수 관례로 재해석해도
우연히 여전히 0x0000-0x03FF 안에 떨어질 수 있어 경고가 안 뜬다. 이건 결함이
아니라 이 휴리스틱의 본질적 한계다 -- magnitude 만으로는 두 관례를 항상
구분할 수 없다.

Ruling 36 -- format_skip_notes 는 skip 구간의 start/end 가 각각 따로
None 일 수 있다는 걸 놓쳐서(특히 xlsx 오염 구간이 파일 끝까지 정상 주소로
안 돌아오는 경우 -- start 는 알지만 end 를 못 정하는 경우) None 을 그대로
'{:04X}' 로 포매팅하려다 TypeError 로 죽었다. 모든 start/end 조합(둘 다
있음/둘 다 없음/한쪽만 없음)을 말로 표현하도록 고쳤다 -- 잘못된 입력
파일이 트레이스백을 내면 안 된다(계측 벤치에서 실행하는 도구이므로).
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from .bench import Bench
from .board.bringup import bring_up_tx, make_chip
from .board.firehawk import FH
from .setup_tx import DEFAULT_CONFIG, _force_utf8_stdout

# xlsx 주소 열 corruption 가드: 정상 파일은 주소가 (거의) 연속 증가한다. 그보다
# 훨씬 큰 폭으로 뛰거나 감소하면 지수표기 오염으로 보고 버린다. 실제 오염 행의
# 델타는 6만 이상이라(0x10000 - 0x10df 등) 여유를 크게 둬도 안전하다.
_XLSX_MAX_ADDR_GAP = 0x100

# 알려진 레지스터 주소 공간(Ruling 35 plausibility check용, 반열림 구간):
# 빔테이블 하위 주소(0x0000-0x03FF) + config/bias 구간(0x1000-0x1204, CLI
# --range 기본값과 동일). _addr_from_xlsx_cell 의 "int 셀=헥사 자릿수" 가정이
# 이 파일에 안 맞으면(예: 평범한 정수 주소 관례) 파싱된 주소 대부분이 이 밖에
# 떨어진다 -- 그 신호를 잡아 경고한다.
_KNOWN_ADDR_RANGES = ((0x0000, 0x0400), (0x1000, 0x1204))
# "significant share"의 기준선. 실측 파일은 알려진 범위를 전혀 벗어나지 않고,
# 관례가 안 맞는 파일은 헥사-자릿수 해석 특성상 대개 대부분(보통 전부) 벗어나므로
# 낮게 잡아도 오탐 없이 잡힌다(정상 덤프에 소수 미지 레지스터가 섞여도 안전).
_XLSX_ADDR_PLAUSIBILITY_THRESHOLD = 0.2


def _looks_like_known_register(addr):
    return any(lo <= addr < hi for lo, hi in _KNOWN_ADDR_RANGES)


class RegDump(dict):
    """{addr: val} 결과 + 파싱에 실패해 건너뛴 행 정보(Ruling 34).

    일반 dict 와 완전히 동등하게 비교/사용된다(``RegDump(...) == {...}`` 가
    성립하므로 기존 호출부·테스트를 안 건드려도 된다) -- 건너뛴 행은 부가
    속성 ``.skipped`` 로만 노출한다.

    ``skipped`` 는 ``{"start": int|None, "end": int|None, "count": int}``
    의 리스트다. 주소를 알 수 있으면(값 셀만 깨진 경우) start==end 인 1개
    구간으로, 주소 자체를 복구할 수 없으면(xlsx 지수표기 오염처럼) 앞뒤 정상
    주소로부터 추정한 구간으로, 그마저 불가능하면(일반 텍스트 포맷) start/end
    가 둘 다 None 인 채로 count 만 담는다.
    """

    def __init__(self, *args, skipped=None, **kwargs):
        super().__init__(*args, **kwargs)
        self.skipped = list(skipped) if skipped else []


def format_skip_notes(skipped):
    """RegDump.skipped -> ["N row(s) [at 0xAAAA[-0xBBBB]] could not be parsed and were not compared", ...].

    Ruling 36: a run's start/end can each independently be known or unknown
    (e.g. an unresolved xlsx corruption run that never finds a known-good
    address before end-of-file has a known start but no end) -- every
    combination is handled in words instead of ever formatting a None as
    hex (that crashed with "unsupported format string passed to
    NoneType.__format__" before this fix).
    """
    notes = []
    for run in skipped:
        n = run["count"]
        start, end = run["start"], run["end"]
        if start is None and end is None:
            where = ""
        elif start is None:
            where = f" ending at 0x{end:04X} (start address undetermined)"
        elif end is None:
            where = f" starting at 0x{start:04X} (end address undetermined -- file ended before a known-good address resumed)"
        elif start == end:
            where = f" at 0x{start:04X}"
        else:
            where = f" at 0x{start:04X}-0x{end:04X}"
        notes.append(f"{n} row(s){where} could not be parsed and were not compared")
    return notes


def _coalesce_addrs(addrs):
    """정렬 가능한 주소 리스트 -> 연속 구간 [{"start","end","count"}, ...]."""
    runs = []
    for a in sorted(addrs):
        if runs and a == runs[-1]["end"] + 1:
            runs[-1]["end"] = a
            runs[-1]["count"] += 1
        else:
            runs.append({"start": a, "end": a, "count": 1})
    return runs


def _to_int(x):
    """'0x1000' / '4096' / 4096 -> int. 변환 불가면 None."""
    if isinstance(x, int):
        return x
    if x is None:
        return None
    s = str(x).strip()
    if not s:
        return None
    try:
        return int(s, 0)
    except ValueError:
        return None


def _addr_from_xlsx_cell(cell):
    """xlsx 주소 셀 -> int, 항상 16진수로 해석한다. 변환 불가면 None.

    실제 Sivers 파일 규칙: int 셀은 헥사 자릿수를 그대로 십진수로 적은 것(예:
    1004 -> 0x1004), 문자열 셀은 '0x' 접두 유무에 관계없이 헥사 문자열이다
    (예: '100A', '0x1200'). 헤더 행("Register Adress" 등 순수 텍스트)은
    16진수로도 파싱되지 않으므로 여기서 자연스럽게 None 이 되어 걸러진다.

    ** 주의(Ruling 35): 이 "int 셀 = 헥사 자릿수" 규칙은 reference/*.xlsx
    실물에서 확인한 이 파일만의 관례다. 평범한 xlsx(정수 셀이 진짜 그 정수
    주소값을 뜻하는 파일)에 이 함수를 그대로 쓰면 주소가 전부 조용히 틀리게
    읽힌다 -- 에러 없이, 그럴듯한 "N differing / M compared" 한 줄만 남긴다.
    자동 판별/전환은 하지 않는다(잘못 추측해 조용히 바꿔치기하는 것도 똑같이
    위험하다) -- 대신 _load_xlsx 끝에서 파싱된 주소가 알려진 레지스터 범위
    밖으로 얼마나 벗어났는지 보고 경고한다(_KNOWN_ADDR_RANGES 참고). **
    """
    if isinstance(cell, int):
        try:
            return int(str(cell), 16)
        except ValueError:
            return None
    if isinstance(cell, str):
        s = cell.strip()
        if not s:
            return None
        try:
            return int(s, 16)
        except ValueError:
            return None
    return None


def _load_xlsx(path, base):
    """xlsx 첫 시트에서 (addr, value) 두 열을 읽는다.

    실제 파일은 데이터가 열 B/C에 있고 열 A가 항상 비어 있으므로, 행의 앞 두
    칸이 아니라 행 전체에서 비어 있지 않은 첫 두 셀을 취한다. 헤더 행은
    자연히 걸러진다(주소 칸이 16진수로도 파싱되지 않으므로). 값만 있는 행은
    base 부터 순차 주소를 매긴다(MATLAB mem_dump 형식 대응).

    파싱 실패 행 두 가지를 구분해서 기록한다(Ruling 34):
      - 주소는 정상(직전 주소에서 이어짐)인데 값 칸이 깨진 경우 -> 그 정확한
        주소를 안다.
      - 주소 칸 자체가 직전 주소보다 작거나 너무 멀리 뛴 경우(Excel 지수표기
        오염 등) -> 그 행의 진짜 주소는 복구 불가능하지만, 앞뒤로 확실히 아는
        정상 주소 사이의 간격으로 범위를 추정할 수 있다.
    """
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    out, seq, prev_addr = {}, base, None
    known_skip_addrs = []
    skip_runs = []
    unknown_run_start = None
    unknown_run_count = 0

    def flush_unknown_run(next_addr):
        nonlocal unknown_run_count
        if unknown_run_count:
            start = (unknown_run_start + 1) if unknown_run_start is not None else None
            end = (next_addr - 1) if next_addr is not None else None
            skip_runs.append({"start": start, "end": end, "count": unknown_run_count})
            unknown_run_count = 0

    for row in ws.iter_rows(values_only=True):
        if not row:
            continue
        cells = [c for c in row if c is not None and str(c).strip() != ""][:2]
        if not cells:
            continue
        if len(cells) < 2:
            val = _to_int(cells[0])
            if val is None:
                continue
            out[seq] = val & 0xFFFF
            seq += 1
            continue
        addr = _addr_from_xlsx_cell(cells[0])
        if addr is None:
            continue  # header text ("Register Adress"/"Value") -- not a data row
        gap_bad = prev_addr is not None and (addr <= prev_addr or addr - prev_addr > _XLSX_MAX_ADDR_GAP)
        if gap_bad:
            if unknown_run_count == 0:
                unknown_run_start = prev_addr
            unknown_run_count += 1
            continue
        flush_unknown_run(addr)
        val = _to_int(cells[1])
        if val is None:
            known_skip_addrs.append(addr)
            prev_addr = addr  # address is known-good; keep the gap guard anchored on it
            continue
        out[addr] = val & 0xFFFF
        prev_addr = addr
    flush_unknown_run(None)
    wb.close()

    skip_runs = _coalesce_addrs(known_skip_addrs) + skip_runs
    skip_runs.sort(key=lambda r: (r["start"] is None, r["start"]))
    for note in format_skip_notes(skip_runs):
        print(f"  note: {note}")

    if out:
        n_total = len(out)
        n_outside = sum(1 for a in out if not _looks_like_known_register(a))
        if n_outside / n_total > _XLSX_ADDR_PLAUSIBILITY_THRESHOLD:
            print(
                "  warning: addresses were read assuming hex-digits-as-decimal "
                "(cell 1004 -> 0x1004); "
                f"{n_outside} of {n_total} parsed addresses fall outside the "
                "known register range (0x0000-0x03FF, 0x1000-0x1203), so this "
                "file may use plain decimal addresses instead."
            )

    return RegDump(out, skipped=skip_runs)


def load_dump(src, base=0x1000):
    """레지스터 덤프를 RegDump({addr: val} + .skipped) 로 로드한다.

    .csv / .xlsx / dict 를 받는다. RegDump 는 dict 서브클래스라 기존처럼
    ``load_dump(...) == {addr: val, ...}`` 비교가 그대로 성립한다.

    csv 허용 형식(자동 감지): 'addr,value' / 'addr value' / 'addr\\tvalue'.
    값만 한 열로 나열된 경우는 base 부터 순차 주소를 매긴다.
    """
    if isinstance(src, dict):
        return RegDump({int(k): int(v) & 0xFFFF for k, v in src.items()})
    p = Path(src)
    if p.suffix.lower() in (".xlsx", ".xlsm"):
        return _load_xlsx(p, base)
    out, seq = {}, base
    known_skip_addrs = []
    unknown_skip_count = 0
    for raw in p.read_text(encoding="utf-8").splitlines():
        s = raw.strip()
        if not s or s[0] in "#%" or s.lower().startswith("addr"):
            continue
        parts = s.replace(",", " ").replace("\t", " ").split()
        nums = [_to_int(x) for x in parts]
        if any(n is None for n in nums):
            if nums and nums[0] is not None:
                known_skip_addrs.append(nums[0])
            else:
                unknown_skip_count += 1
            continue
        if len(nums) >= 2:
            out[nums[0]] = nums[1] & 0xFFFF
        elif len(nums) == 1:
            out[seq] = nums[0] & 0xFFFF
            seq += 1
    skip_runs = _coalesce_addrs(known_skip_addrs)
    if unknown_skip_count:
        skip_runs.append({"start": None, "end": None, "count": unknown_skip_count})
    for note in format_skip_notes(skip_runs):
        print(f"  note: {note}")
    return RegDump(out, skipped=skip_runs)


def save_dump(d, path):
    """{addr: val} 을 'addr,value' hex CSV 로 저장한다(주소 오름차순)."""
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    with p.open("w", encoding="utf-8") as f:
        f.write("addr,value\n")
        for addr in sorted(d):
            f.write(f"0x{addr:04X},0x{d[addr] & 0xFFFF:04X}\n")
    print(f"  dumped {len(d)} regs -> {p}")


def regdiff(theirs, ours, log=print):
    """두 덤프를 비교해 다른 레지스터만 출력하고, 다른 주소 목록을 반환한다."""
    addrs = sorted(set(theirs) | set(ours))
    diffs = []
    for a in addrs:
        tv, ov = theirs.get(a), ours.get(a)
        if tv != ov:
            diffs.append(a)
            ts = f"0x{tv:04X}" if tv is not None else "  --  "
            os_ = f"0x{ov:04X}" if ov is not None else "  --  "
            log(f"  0x{a:04X}:  theirs={ts}  ours={os_}")
    log(f"  {len(diffs)} differing / {len(addrs)} compared")
    return diffs


def main(argv=None):
    _force_utf8_stdout()
    ap = argparse.ArgumentParser(description="Dump and compare chip registers.")
    ap.add_argument("--fake", action="store_true", help="no hardware (fake SPI)")
    ap.add_argument("--no-power", action="store_true", help="skip PSU power-up (already powered)")
    ap.add_argument("--save", metavar="PATH", help="save the dump as addr,value CSV")
    ap.add_argument("--diff", metavar="PATH", help="reference dump to compare against (.csv/.xlsx)")
    ap.add_argument("--ours", metavar="PATH", help="our dump from a file (default: read the chip live)")
    ap.add_argument("--range", nargs=2, metavar=("A", "B"), default=["0x1000", "0x1204"],
                     help="dump address range [A, B) (default 0x1000 0x1204)")
    ap.add_argument("--config", type=Path, default=DEFAULT_CONFIG, help="bench toml path")
    args = ap.parse_args(argv)

    a, b = int(args.range[0], 0), int(args.range[1], 0)

    if args.ours:
        ours = load_dump(args.ours)
    else:
        bench = Bench.from_toml(args.config, fake=args.fake)
        bench.connect_all()
        if not args.no_power:
            bench.power_up()
        else:
            print("[regdump] skipping power-up (--no-power)")
        chip = make_chip(bench.board, fake=args.fake)
        bring_up_tx(chip, bench.board, require_version=not args.fake)
        fh = FH(chip, bench.board.chip_id)
        ours = {addr: fh.rd(addr) for addr in range(a, b)}

    if args.save:
        save_dump(ours, args.save)
    if args.diff:
        theirs = load_dump(args.diff, base=a)
        diffs = regdiff(theirs, ours)
        return 1 if diffs else 0
    if not args.save:
        for addr in sorted(ours):
            print(f"  0x{addr:04X} = 0x{ours[addr]:04X}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
