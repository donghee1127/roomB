# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------- 색상 팔레트 ----------
NAVY = RGBColor(0x1A, 0x1F, 0x2E)
DARK_BG = RGBColor(0x14, 0x17, 0x21)
ORANGE = RGBColor(0xE0, 0x7A, 0x3C)
WHITE = RGBColor(0xF5, 0xF5, 0xF3)
GRAY = RGBColor(0xB8, 0xBD, 0xC7)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
blank_layout = prs.slide_layouts[6]


def add_bg(slide, color=DARK_BG):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    slide.shapes._spTree.remove(shape._element)
    slide.shapes._spTree.insert(2, shape._element)
    return shape


def add_accent_bar(slide, top=True):
    y = Inches(0) if top else SLIDE_H - Inches(0.12)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, y, SLIDE_W, Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ORANGE
    bar.line.fill.background()
    bar.shadow.inherit = False


def add_textbox(slide, left, top, width, height, text, size=18, color=WHITE,
                 bold=False, align=PP_ALIGN.LEFT, font="맑은 고딕", anchor=None):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    if anchor is not None:
        tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    return tb


def add_bullets(slide, left, top, width, height, items, size=20, color=WHITE,
                 font="맑은 고딕", line_spacing=1.3):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = level
        p.line_spacing = line_spacing
        p.space_after = Pt(10)
        bullet = "▸ " if level == 0 else "–  "
        run = p.add_run()
        run.text = bullet + text
        run.font.size = Pt(size - level * 2)
        run.font.color.rgb = color if level == 0 else GRAY
        run.font.name = font
    return tb


def add_slide_title(slide, title, kicker=None):
    if kicker:
        add_textbox(slide, Inches(0.7), Inches(0.35), Inches(10), Inches(0.4),
                    kicker, size=14, color=ORANGE, bold=True)
    add_textbox(slide, Inches(0.7), Inches(0.7), Inches(11.5), Inches(0.9),
                title, size=32, color=WHITE, bold=True)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.55),
                                   Inches(1.2), Pt(3))
    line.fill.solid()
    line.fill.fore_color.rgb = ORANGE
    line.line.fill.background()
    line.shadow.inherit = False


def new_slide():
    slide = prs.slides.add_slide(blank_layout)
    add_bg(slide)
    add_accent_bar(slide, top=False)
    return slide


# ---------- 1. 표지 ----------
slide = new_slide()
add_bg(slide, NAVY)
add_accent_bar(slide, top=True)
add_accent_bar(slide, top=False)
add_textbox(slide, Inches(1), Inches(2.6), Inches(11.3), Inches(1.4),
            "Claude Code 사용법", size=54, color=WHITE, bold=True, align=PP_ALIGN.LEFT)
add_textbox(slide, Inches(1), Inches(3.75), Inches(11.3), Inches(0.7),
            "터미널에서 바로 쓰는 AI 코딩 어시스턴트", size=22, color=ORANGE)
add_textbox(slide, Inches(1), Inches(6.6), Inches(6), Inches(0.5),
            "2026", size=16, color=GRAY)

# ---------- 2. Claude Code란? ----------
slide = new_slide()
add_slide_title(slide, "Claude Code란?", kicker="INTRO")
add_bullets(slide, Inches(0.9), Inches(2.0), Inches(11.5), Inches(4.8), [
    "Anthropic이 만든 터미널(CLI) 기반 AI 코딩 에이전트",
    "자연어로 요청하면 코드 읽기·작성·수정·실행을 직접 수행",
    ("파일 검색/편집, 셸 명령 실행, Git 작업까지 하나의 대화로 처리", 1),
    "VS Code, JetBrains 등 IDE 확장과 데스크톱/웹 앱에서도 사용 가능",
    "복잡한 리팩터링부터 버그 수정, 문서화, 테스트 작성까지 폭넓게 활용",
])

# ---------- 3. 설치 방법 ----------
slide = new_slide()
add_slide_title(slide, "설치 방법", kicker="GETTING STARTED")
add_bullets(slide, Inches(0.9), Inches(2.0), Inches(11.5), Inches(2.0), [
    "Node.js 설치 후 npm으로 전역 설치",
    "터미널에서 아래 명령 실행:",
])
code_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(3.3),
                                   Inches(8), Inches(0.9))
code_box.fill.solid()
code_box.fill.fore_color.rgb = RGBColor(0x22, 0x27, 0x34)
code_box.line.color.rgb = ORANGE
code_box.line.width = Pt(1)
code_box.shadow.inherit = False
tf = code_box.text_frame
tf.word_wrap = True
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.LEFT
run = p.add_run()
run.text = "  npm install -g @anthropic-ai/claude-code"
run.font.name = "Consolas"
run.font.size = Pt(18)
run.font.color.rgb = ORANGE
add_bullets(slide, Inches(0.9), Inches(4.6), Inches(11.5), Inches(2.2), [
    "설치 후 원하는 프로젝트 폴더에서 claude 명령으로 실행",
    "최초 실행 시 로그인/인증 절차 진행",
])

# ---------- 4. 시작하기 ----------
slide = new_slide()
add_slide_title(slide, "시작하기", kicker="GETTING STARTED")
add_bullets(slide, Inches(0.9), Inches(2.0), Inches(11.5), Inches(4.8), [
    "프로젝트 디렉터리로 이동 후 claude 입력",
    "자연어로 작업 요청 (예: \"로그인 버그 고쳐줘\", \"테스트 추가해줘\")",
    ("Claude가 관련 파일을 스스로 탐색하고 계획을 세워 진행", 1),
    "코드 수정·명령 실행 전, 필요 시 승인(permission) 요청",
    ("위험하거나 되돌리기 어려운 작업은 항상 확인을 거침", 1),
    "대화는 세션 내내 이어지며 문맥을 기억",
])

# ---------- 5. 핵심 기능 ----------
slide = new_slide()
add_slide_title(slide, "핵심 기능", kicker="FEATURES")
add_bullets(slide, Inches(0.9), Inches(2.0), Inches(11.5), Inches(4.8), [
    "파일 읽기·검색 — Read, Glob, Grep으로 코드베이스 탐색",
    "파일 편집·생성 — Edit, Write로 정확한 코드 변경",
    "셸 명령 실행 — Bash/PowerShell로 빌드, 테스트, 스크립트 실행",
    "Git 연동 — 커밋, PR 생성, 브랜치 관리 등 지원",
    "웹 검색·문서 열람 — 최신 정보나 API 문서 참고 가능",
    "이미지·PDF 인식 — 스크린샷, 문서 파일 분석",
])

# ---------- 6. 슬래시 명령어 ----------
slide = new_slide()
add_slide_title(slide, "주요 슬래시 명령어", kicker="COMMANDS")
add_bullets(slide, Inches(0.9), Inches(2.0), Inches(11.5), Inches(4.8), [
    "/help — 사용법 및 도움말 확인",
    "/init — 프로젝트 문서(CLAUDE.md) 초기 생성",
    "/config — 모델, 테마 등 설정 변경",
    "/clear — 대화 컨텍스트 초기화",
    "/agents — 서브에이전트 관리",
    "/review 또는 /code-review — 코드 리뷰 실행",
    "사용자 정의 슬래시 명령도 프로젝트에 추가 가능",
])

# ---------- 7. Plan Mode ----------
slide = new_slide()
add_slide_title(slide, "Plan Mode", kicker="WORKFLOW")
add_bullets(slide, Inches(0.9), Inches(2.0), Inches(11.5), Inches(4.8), [
    "코드를 바로 수정하지 않고 먼저 계획을 세우는 모드",
    "탐색 → 설계 → 계획서 작성 → 사용자 승인 순으로 진행",
    ("승인 전까지는 읽기 전용 작업만 수행", 1),
    "복잡하거나 영향 범위가 큰 작업에 유용",
    "승인 후에는 계획에 따라 실제 구현 진행",
])

# ---------- 8. Sub-agents & Skills ----------
slide = new_slide()
add_slide_title(slide, "서브에이전트 & 스킬", kicker="ADVANCED")
add_bullets(slide, Inches(0.9), Inches(2.0), Inches(11.5), Inches(4.8), [
    "서브에이전트(Sub-agent) — 특정 작업에 특화된 독립 에이전트 실행",
    ("탐색 전용(Explore), 코드 리뷰 등 목적별로 활용", 1),
    "Skill — 반복되는 작업 절차를 패키징해 재사용",
    ("예: 배포 체크리스트, 코드 리뷰 규칙, 문서 초기화 등", 1),
    "여러 에이전트를 병렬로 실행해 대규모 작업을 빠르게 처리 가능",
])

# ---------- 9. MCP & 확장 ----------
slide = new_slide()
add_slide_title(slide, "MCP & 확장", kicker="INTEGRATIONS")
add_bullets(slide, Inches(0.9), Inches(2.0), Inches(11.5), Inches(4.8), [
    "MCP(Model Context Protocol) — 외부 도구·서비스 연동 표준",
    ("Slack, 데이터베이스, 사내 시스템 등과 연결 가능", 1),
    "Hooks — 특정 이벤트(도구 실행 전/후 등)에 자동 스크립트 실행",
    "설정 파일(settings.json)로 권한, 훅, 환경변수 등 세밀하게 제어",
])

# ---------- 10. 실전 팁 ----------
slide = new_slide()
add_slide_title(slide, "실전 팁", kicker="BEST PRACTICES")
add_bullets(slide, Inches(0.9), Inches(2.0), Inches(11.5), Inches(4.8), [
    "요청은 구체적으로 — 목표, 제약조건, 원하는 결과를 명확히 전달",
    "큰 작업은 Plan Mode로 먼저 설계한 뒤 진행",
    "CLAUDE.md에 프로젝트 규칙·컨벤션을 기록해 일관성 유지",
    "위험한 작업(삭제, 강제 push 등)은 항상 검토 후 승인",
    "반복 작업은 커스텀 명령어나 훅으로 자동화",
])

# ---------- 11. 마무리 ----------
slide = new_slide()
add_bg(slide, NAVY)
add_accent_bar(slide, top=True)
add_accent_bar(slide, top=False)
add_textbox(slide, Inches(1), Inches(2.6), Inches(11.3), Inches(1.0),
            "감사합니다", size=44, color=WHITE, bold=True)
add_textbox(slide, Inches(1), Inches(3.6), Inches(11.3), Inches(0.6),
            "더 알아보기: Claude Code 공식 문서 (docs.claude.com)", size=18, color=GRAY)

prs.save("Claude_Code_사용법.pptx")
print("Saved: Claude_Code_사용법.pptx")
